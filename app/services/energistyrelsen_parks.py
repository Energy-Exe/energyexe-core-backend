"""Parkproduktion: park-metered production split equally over the park's turbines.

Since 2026 the Vinddatasæt leaves turbines that are metered collectively
blank and publishes their park total in a second workbook. The legacy
import represented such parks as an equal split over the park's turbines
(verified against the database), so the same rule is applied here.
"""

from collections import defaultdict
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO

from openpyxl import load_workbook

from app.services.energistyrelsen_workbook import _month, _number


def parse_parkproduktion(content: bytes, start_month: str, end_month: str):
    """Return {(park GUID, YYYY-MM): kWh} for the requested months (duplicate rows summed)."""
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet_name = next((name for name in book.sheetnames if "produktion" in name.lower()), None)
    if sheet_name is None:
        raise ValueError("Expected worksheet 'Vindmølleparkers produktion'")
    rows = book[sheet_name].iter_rows(values_only=True)
    headers = None
    row_number = 0
    for row_number, row in enumerate(rows, 1):
        if row and row[0] is not None and "parknummer" in str(row[0]).lower():
            headers = row
            break
    if headers is None:
        raise ValueError("Expected 'Parknummer' header in column A")
    months = {i: month for i, heading in enumerate(headers) if (month := _month(heading)) and start_month <= month <= end_month}
    if not months:
        raise ValueError("No monthly columns in requested range")
    values = {}
    for row_number, row in enumerate(rows, row_number + 1):
        park = str(row[0]).strip().upper() if row and row[0] is not None else ""
        if not park:
            continue
        for col, month in months.items():
            value = _number(row[col] if col < len(row) else None)
            if value is None:
                continue
            key = (park, month)
            # A park re-registered mid-month appears twice; both rows are
            # partial readings of the same park, so they add up.
            values[key] = values.get(key, Decimal(0)) + value
    return values


def month_range(month):
    year, number = map(int, month.split("-"))
    return date(year, number, 1), date(year + (number == 12), number % 12 + 1, 1)


def is_active(attributes, month):
    """A turbine counts for a month when connected before the month ends
    and not decommissioned before it starts."""
    start, end = month_range(month)
    connected = attributes.get("connected")
    decommissioned = attributes.get("decommissioned")
    return (connected is None or connected < end) and (decommissioned is None or decommissioned >= start)


def split_park_values(park_values, attributes, matched, values):
    """Add per-turbine shares of each park total to ``values`` in place.

    ``attributes`` is {gsrn: turbine_attributes(...)}, ``matched`` the GSRNs
    that exist in turbine_units. A turbine's own Vinddatasæt reading always
    wins over its park share. Returns (provenance, report).
    """
    turbines_by_park = defaultdict(list)
    for gsrn, attrs in attributes.items():
        if attrs.get("park"):
            turbines_by_park[attrs["park"]].append(gsrn)
    provenance = {}
    report = {"mapped": {}, "unmapped_with_production": {}, "no_active_turbines": [], "direct_wins": 0}
    by_park = defaultdict(dict)
    for (park, month), kwh in park_values.items():
        by_park[park][month] = kwh
    for park, months in by_park.items():
        gsrns = turbines_by_park.get(park, [])
        park_matched = [gsrn for gsrn in gsrns if gsrn in matched]
        if not park_matched:
            report["unmapped_with_production"][park] = {
                "n_turbines_in_file": len(gsrns),
                "first_month": min(months), "last_month": max(months), "n_months": len(months),
            }
            continue
        written = 0
        for month, kwh in sorted(months.items()):
            active = [gsrn for gsrn in gsrns if is_active(attributes[gsrn], month)]
            if not active:
                report["no_active_turbines"].append([park, month])
                continue
            share = (kwh / 1000 / len(active)).quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)
            for gsrn in active:
                if gsrn not in matched:
                    continue
                key = (gsrn, month)
                if key in values:
                    report["direct_wins"] += 1
                    continue
                values[key] = share
                provenance[key] = {"park": park, "split_n": len(active), "park_kwh": str(kwh)}
                written += 1
        report["mapped"][park] = {"n_turbines_in_file": len(gsrns), "n_matched": len(park_matched),
                                  "months_with_value": len(months), "values_written": written}
    return provenance, report
